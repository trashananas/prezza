# Модуль для парсинга презентаций (PDF/PPTX)
import io
import fitz  # PyMuPDF
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def parse_presentation(uploaded_file):
    filename = uploaded_file.name.lower()
    slides = []
    if filename.endswith('.pdf'):
        pdf_bytes = uploaded_file.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        for i, page in enumerate(doc, 1):
            text = page.get_text().strip()

            # Попытка извлечения таблиц
            tables_content = ""
            try:
                tabs = page.find_tables()
                for tab in tabs:
                    df = tab.to_pandas()
                    if not df.empty:
                        tables_content += "\n\n[Таблица с данными]:\n" + df.to_markdown(index=False)
            except Exception as e:
                tables_content += f"\n\n[Ошибка при извлечении таблицы: {e}]"

            # Попытка извлечения изображений
            image_list = page.get_images(full=True)
            images_found = len(image_list)
            if images_found > 0:
                text += f"\n\n[На слайде найдено изображений: {images_found}]"

            slides.append({
                "title": f"Слайд {i}",
                "content": text + tables_content,
                "images": []
            })
    elif filename.endswith('.pptx'):
        pptx_bytes = uploaded_file.read()
        prs = Presentation(io.BytesIO(pptx_bytes))
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            tables_content = ""
            images_count = 0
            charts_content = ""

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())

                if shape.has_table:
                    try:
                        table = shape.table
                        rows = []
                        for row in table.rows:
                            rows.append([cell.text_frame.text.strip() for cell in row.cells])

                        if rows:
                            header = rows[0]
                            body = rows[1:]
                            table_md = "| " + " | ".join(header) + " |\n"
                            table_md += "| " + " | ".join(["---"] * len(header)) + " |\n"
                            for row in body:
                                table_md += "| " + " | ".join(row) + " |\n"
                            tables_content += "\n\n[Таблица с данными]:\n" + table_md
                    except Exception as e:
                        tables_content += f"\n\n[Ошибка при парсинге таблицы: {e}]"

                if shape.has_chart:
                    try:
                        chart = shape.chart
                        title = chart.chart_title.text_frame.text if chart.has_title else 'Без названия'
                        charts_content += f"\n\n[Диаграмма: {title}]\n"
                        for series in chart.plots[0].series:
                            charts_content += f"- Серия '{series.name}': {list(series.values)}\n"
                    except Exception as e:
                        charts_content += f"\n[Ошибка при извлечении данных диаграммы: {e}]\n"

                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    images_count += 1

            slide_text = "\n".join(texts).strip()
            if images_count > 0:
                slide_text += f"\n\n[На слайде найдено изображений: {images_count}]"

            slides.append({
                "title": f"Слайд {i}",
                "content": slide_text + tables_content + charts_content,
                "images": []
            })
    else:
        slides.append({"title": "Ошибка", "content": "Неподдерживаемый формат файла", "images": []})
    return slides
